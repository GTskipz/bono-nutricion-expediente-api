from __future__ import annotations

from typing import Optional, List, Tuple
from datetime import datetime
from openpyxl import Workbook, load_workbook
from io import BytesIO
import threading

import random
import string

from fastapi import HTTPException
from sqlalchemy.orm import Session
from sqlalchemy import func, or_, case

from app.models.expediente_electronico import ExpedienteElectronico
from app.models.cat_departamento import CatDepartamento
from app.models.cat_municipio import CatMunicipio
from app.models.cat_estado_flujo_expediente import CatEstadoFlujoExpediente

from app.models.cuentas_bancarias import (
    LoteAperturaCuenta,
    DetalleAperturaCuenta,
    CuentaBancariaExpediente,
)

from app.services.banco_archivo_service import upload_archivo_banco_core
from app.services.tracking_evento_service import TrackingEventoService


# =========================================================
# BANDEJA (expedientes por estado_flujo_id)
# =========================================================
def bandeja_expedientes_por_estado_flujo(
    db: Session,
    *,
    estado_flujo_id: int,
    texto: Optional[str] = None,
    departamento_id: Optional[int] = None,
    municipio_id: Optional[int] = None,
    page: int = 1,
    limit: int = 20,
):

    texto = (texto or "").strip()
    filters = [ExpedienteElectronico.estado_flujo_id == estado_flujo_id]

    if texto:
        filters.append(
            or_(
                ExpedienteElectronico.nombre_beneficiario.ilike(f"%{texto}%"),
                ExpedienteElectronico.cui_beneficiario.like(f"{texto}%"),
            )
        )

    if departamento_id:
        filters.append(ExpedienteElectronico.departamento_id == departamento_id)

    if municipio_id:
        filters.append(ExpedienteElectronico.municipio_id == municipio_id)

    total = db.query(func.count(ExpedienteElectronico.id)).filter(*filters).scalar() or 0

    offset = (page - 1) * limit

    rows = (
        db.query(
            ExpedienteElectronico.id,
            ExpedienteElectronico.created_at,
            ExpedienteElectronico.nombre_beneficiario,
            ExpedienteElectronico.cui_beneficiario,
            ExpedienteElectronico.estado_expediente,
            ExpedienteElectronico.bpm_status,
            ExpedienteElectronico.bpm_current_task_name,
            CatEstadoFlujoExpediente.codigo.label("estado_flujo_codigo"),
            CatEstadoFlujoExpediente.nombre.label("estado_flujo_nombre"),
            CatDepartamento.nombre.label("departamento"),
            CatMunicipio.nombre.label("municipio"),
        )
        .outerjoin(
            CatEstadoFlujoExpediente,
            CatEstadoFlujoExpediente.id == ExpedienteElectronico.estado_flujo_id,
        )
        .outerjoin(CatDepartamento, CatDepartamento.id == ExpedienteElectronico.departamento_id)
        .outerjoin(CatMunicipio, CatMunicipio.id == ExpedienteElectronico.municipio_id)
        .filter(*filters)
        .order_by(ExpedienteElectronico.created_at.desc())
        .offset(offset)
        .limit(limit)
        .all()
    )

    data = []
    for r in rows:
        data.append({
            "id": r.id,
            "created_at": r.created_at,
            "nombre_beneficiario": r.nombre_beneficiario,
            "cui_beneficiario": r.cui_beneficiario,
            "estado_expediente": r.estado_expediente,
            "bpm_status": r.bpm_status,
            "bpm_current_task_name": r.bpm_current_task_name,
            "estado_flujo_codigo": getattr(r, "estado_flujo_codigo", None),
            "estado_flujo_nombre": getattr(r, "estado_flujo_nombre", None),
            "departamento": r.departamento,
            "municipio": r.municipio,
        })

    return {"data": data, "page": page, "limit": limit, "total": total}


# =========================================================
# CREAR LOTE
# =========================================================
def crear_lote_apertura(
    db: Session,
    *,
    expediente_ids: List[int],
    creado_por: Optional[str] = None,
    observacion: Optional[str] = None,
    proveedor_servicio: Optional[str] = None,
) -> Tuple[int, int]:

    exp_rows = (
        db.query(
            ExpedienteElectronico.id,
            ExpedienteElectronico.cui_beneficiario,
            ExpedienteElectronico.nombre_beneficiario,
            CatDepartamento.nombre.label("departamento"),
            CatMunicipio.nombre.label("municipio"),
        )
        .outerjoin(CatDepartamento, CatDepartamento.id == ExpedienteElectronico.departamento_id)
        .outerjoin(CatMunicipio, CatMunicipio.id == ExpedienteElectronico.municipio_id)
        .filter(ExpedienteElectronico.id.in_(expediente_ids))
        .all()
    )

    lote = LoteAperturaCuenta(
        banco_codigo="BANRURAL",
        estado="CREADO",
        creado_por=creado_por,
        observacion=observacion,
        proveedor_servicio=proveedor_servicio,
    )

    db.add(lote)
    db.flush()

    for r in exp_rows:
        item = DetalleAperturaCuenta(
            lote_id=lote.id,
            expediente_id=r.id,
            estado="PENDIENTE",
            cui_beneficiario=r.cui_beneficiario,
            nombre_beneficiario=r.nombre_beneficiario,
            titular_dpi=r.cui_beneficiario,
            titular_nombre=r.nombre_beneficiario,
            departamento=r.departamento,
            municipio=r.municipio,
        )
        db.add(item)

    db.commit()

    return lote.id, len(exp_rows)


# =========================================================
# LISTAR LOTES
# =========================================================
def listar_lotes_apertura(
    db: Session,
    *,
    anio: Optional[int] = None,
    estado: Optional[str] = None,
    texto: Optional[str] = None,
    page: int = 1,
    limit: int = 20,
):

    texto = (texto or "").strip()
    offset = (page - 1) * limit

    q = db.query(LoteAperturaCuenta)

    if estado:
        q = q.filter(LoteAperturaCuenta.estado == estado)

    total = q.with_entities(func.count(LoteAperturaCuenta.id)).scalar() or 0

    lotes = q.offset(offset).limit(limit).all()

    data = []
    for l in lotes:
        data.append({
            "id": l.id,
            "estado": l.estado,
            "banco_codigo": l.banco_codigo,
            "creado_en": l.creado_en,
        })

    return {"data": data, "page": page, "limit": limit, "total": total}


# =========================================================
# OBTENER LOTE
# =========================================================
def obtener_lote_apertura(db: Session, *, lote_id: int):

    lote = db.query(LoteAperturaCuenta).filter(LoteAperturaCuenta.id == lote_id).first()

    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    return lote


# =========================================================
# LISTAR ITEMS
# =========================================================
def listar_items_lote(
    db: Session,
    *,
    lote_id: int,
    page: int = 1,
    limit: int = 50,
):

    base_q = db.query(DetalleAperturaCuenta).filter(DetalleAperturaCuenta.lote_id == lote_id)

    total = base_q.with_entities(func.count(DetalleAperturaCuenta.id)).scalar()

    offset = (page - 1) * limit

    rows = base_q.offset(offset).limit(limit).all()

    return {"data": rows, "page": page, "limit": limit, "total": total}


# =========================================================
# LEGACY PROCESAR LOTE
# =========================================================
def procesar_lote_apertura_simulado(db: Session, *, lote_id: int) -> dict:

    total_items = db.query(func.count(DetalleAperturaCuenta.id)).filter(
        DetalleAperturaCuenta.lote_id == lote_id
    ).scalar()

    return {
        "lote_id": lote_id,
        "total_items": total_items,
        "cuentas_creadas": 0,
        "rechazados": 0,
    }


# =========================================================
# PROCESAR RESPUESTA BANCO
# =========================================================
def procesar_respuesta_banco_excel(
    db: Session,
    *,
    lote_id: int,
    file_bytes: bytes,
    filename: str,
):

    lote = (
        db.query(LoteAperturaCuenta)
        .filter(LoteAperturaCuenta.id == lote_id)
        .first()
    )

    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    # =====================================================
    # Guardar archivo en almacenamiento (MinIO)
    # =====================================================

    archivo = upload_archivo_banco_core(
        db=db,
        tipo_operacion="APERTURA_CUENTA",
        operacion_id=lote_id,
        tipo_archivo="RESPUESTA",
        banco_codigo="BANRURAL",
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content=file_bytes,
    )

    # guardar relación lote -> archivo
    lote.archivo_respuesta_id = archivo["id"]

    # =====================================================
    # Procesar Excel
    # =====================================================

    wb = load_workbook(BytesIO(file_bytes))
    ws = wb.active

    cuentas_creadas = 0
    rechazados = 0
    now = datetime.utcnow()

    estado_ok = (
        db.query(CatEstadoFlujoExpediente)
        .filter(CatEstadoFlujoExpediente.codigo == "CUENTA_BANCARIA_CREADA")
        .first()
    )

    for row in ws.iter_rows(min_row=2, values_only=True):

        dpi = str(row[10]).strip() if row[10] else None
        cuenta = str(row[13]).strip() if row[13] else None

        if not dpi:
            continue

        item = (
            db.query(DetalleAperturaCuenta)
            .filter(
                DetalleAperturaCuenta.lote_id == lote_id,
                DetalleAperturaCuenta.titular_dpi == dpi,
            )
            .first()
        )

        if not item:
            continue

        item.procesado_en = now

        if cuenta:

            item.estado = "CUENTA_CREADA"
            item.numero_cuenta = cuenta

            db.add(
                CuentaBancariaExpediente(
                    expediente_id=item.expediente_id,
                    banco_codigo="BANRURAL",
                    numero_cuenta=cuenta,
                    titular_dpi=item.titular_dpi,
                    titular_nombre=item.titular_nombre,
                    detalle_apertura_id=item.id,
                )
            )

            if estado_ok:
                db.query(ExpedienteElectronico).filter(
                    ExpedienteElectronico.id == item.expediente_id
                ).update({"estado_flujo_id": estado_ok.id})

            cuentas_creadas += 1

        else:

            item.estado = "RECHAZADO"
            rechazados += 1

    # =====================================================
    # Actualizar lote
    # =====================================================

    lote.estado = "PROCESADO"
    lote.procesado_en = now

    db.commit()

    # =====================================================
    # Hilo BPM (placeholder)
    # =====================================================

    threading.Thread(
        target=_spiff_placeholder,
        args=(lote_id,),
        daemon=True,
    ).start()

    return {
        "lote_id": lote_id,
        "cuentas_creadas": cuentas_creadas,
        "rechazados": rechazados,
    }

def _spiff_placeholder(lote_id: int):
    pass

def generar_excel_lote_export_bytes(db: Session, *, lote_id: int) -> bytes:

    lote = db.query(LoteAperturaCuenta).filter(
        LoteAperturaCuenta.id == lote_id
    ).first()

    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    items = (
        db.query(DetalleAperturaCuenta)
        .filter(DetalleAperturaCuenta.lote_id == lote_id)
        .order_by(DetalleAperturaCuenta.id.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "APERTURA_CUENTAS"

    headers = [
        "No.",
        "Cod Usuario",
        "Cod Integrante",
        "Primer nombre",
        "Segundo nombre",
        "Tercer nombre",
        "Primer apellido",
        "Segundo apellido",
        "Apellido de casada(o)",
        "No. Orden",
        "Numero de Registro",
        "Genero MF",
        "Direccion Domiciliar",
        "CUENTA",
        "DEPTO",
        "MUNICIPIO",
    ]

    ws.append(headers)

    contador = 1

    for item in items:

        partes = (item.titular_nombre or "").strip().split()

        primer_nombre = ""
        segundo_nombre = ""
        tercer_nombre = ""
        primer_apellido = ""
        segundo_apellido = ""

        if len(partes) == 4:
            primer_nombre = partes[0]
            segundo_nombre = partes[1]
            primer_apellido = partes[2]
            segundo_apellido = partes[3]

        elif len(partes) >= 5:
            primer_nombre = partes[0]
            segundo_nombre = partes[1]
            tercer_nombre = partes[2]
            primer_apellido = partes[3]
            segundo_apellido = partes[4]

        elif len(partes) == 3:
            primer_nombre = partes[0]
            segundo_nombre = partes[1]
            primer_apellido = partes[2]

        elif len(partes) == 2:
            primer_nombre = partes[0]
            primer_apellido = partes[1]

        elif len(partes) == 1:
            primer_nombre = partes[0]

        ws.append([
            contador,

            "",  # Cod Usuario
            "",  # Cod Integrante

            primer_nombre,
            segundo_nombre,
            tercer_nombre,

            primer_apellido,
            segundo_apellido,

            "",  # Apellido casada

            "DPI",
            item.titular_dpi or "",

            "",  # genero

            item.direccion or "",
            "",  # cuenta

            item.departamento or "",
            item.municipio or "",
        ])

        contador += 1

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.getvalue()

def validar_excel_respuesta_banco(file_bytes: bytes):

    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception:
        raise HTTPException(
            status_code=400,
            detail="El archivo Excel no es válido o está corrupto."
        )

    ws = wb.active

    if ws.max_row < 2:
        raise HTTPException(
            status_code=400,
            detail="El archivo Excel no contiene registros."
        )

    headers = []

    for cell in ws[1]:
        if cell.value:
            headers.append(str(cell.value).strip().upper())
        else:
            headers.append("")

    columnas_requeridas = [
        "PRIMER NOMBRE",
        "SEGUNDO NOMBRE",
        "TERCER NOMBRE",
        "PRIMER APELLIDO",
        "SEGUNDO APELLIDO",
        "APELLIDO DE CASADA(O)",
        "CUENTA",
    ]

    faltantes = []

    for col in columnas_requeridas:
        if col not in headers:
            faltantes.append(col)

    if faltantes:
        raise HTTPException(
            status_code=400,
            detail=f"El archivo Excel no contiene las columnas requeridas: {', '.join(faltantes)}"
        )

    return True
