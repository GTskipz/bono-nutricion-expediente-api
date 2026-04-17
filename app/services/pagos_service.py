from __future__ import annotations

import io
from tkinter.font import Font
from typing import Any, Dict, Optional, List, Tuple
from datetime import datetime
from uuid import uuid4

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import String, func, or_, case, text

from openpyxl import Workbook
from openpyxl.styles import Font as XLFont
from io import BytesIO

from app.models.expediente_electronico import ExpedienteElectronico
from app.models.cat_departamento import CatDepartamento
from app.models.cat_municipio import CatMunicipio
from app.models.cat_estado_flujo_expediente import CatEstadoFlujoExpediente

from app.models.cuentas_bancarias import CuentaBancariaExpediente
from app.models.info_general import InfoGeneral
from app.models.pagos import (
    LotePago,
    DetallePago,
    CatFiltroPago,
    CatFiltroPagoOpcion,
    ExpedienteCuentaCorriente,
)

# ✅ Tracking por expediente (eventos importantes)
from app.services.tracking_evento_service import TrackingEventoService


# -------------------------------------------------
# FILTROS DISPONIBLES PARA PAGOS
# -------------------------------------------------
def obtener_filtros_pago(db: Session):
    filtros = (
        db.query(CatFiltroPago)
        .filter(CatFiltroPago.activo == True)
        .order_by(CatFiltroPago.orden)
        .all()
    )

    resultado = []

    for f in filtros:
        item = {
            "codigo": f.codigo,
            "nombre": f.nombre,
            "descripcion": f.descripcion,
            "tipo_control": f.tipo_control,
        }

        if f.tipo_control == "SELECT":
            opciones = (
                db.query(CatFiltroPagoOpcion)
                .filter(
                    CatFiltroPagoOpcion.filtro_id == f.id,
                    CatFiltroPagoOpcion.activo == True,
                )
                .order_by(CatFiltroPagoOpcion.orden)
                .all()
            )

            item["opciones"] = [
                {
                    "valor": op.valor,
                    "etiqueta": op.etiqueta,
                }
                for op in opciones
            ]

        resultado.append(item)

    return resultado

# -------------------------------------------------
# CONSTRUIR WHERE DINAMICO
# -------------------------------------------------
def _build_filtros_where(db: Session, filtros: dict):

    filtros_db = (
        db.query(CatFiltroPago)
        .filter(CatFiltroPago.activo == True)
        .all()
    )

    mapa = {f.codigo: f for f in filtros_db}

    condiciones = []
    params = {}

    for codigo, valor in (filtros or {}).items():

        f = mapa.get(codigo)
        if not f:
            continue

        campo = f.campo_sql
        operador = f.operador
        param_name = f"param_{codigo}"

        # convertir números
        if isinstance(valor, str) and valor.isdigit():
            valor = int(valor)

        if isinstance(valor, list) and valor:

            placeholders = []

            for idx, item in enumerate(valor):

                sub_name = f"{param_name}_{idx}"
                placeholders.append(f":{sub_name}")
                params[sub_name] = item

            condiciones.append(
                f"{campo} IN ({', '.join(placeholders)})"
            )

        else:

            condiciones.append(
                f"{campo} {operador} :{param_name}"
            )

            params[param_name] = valor

    return condiciones, params

# -------------------------------------------------
# PREVIEW LOTE DE PAGOS
# -------------------------------------------------
def preview_lote_pago(
    db: Session,
    *,
    monto_por_persona: float,
    presupuesto_total: float,
    filtros: dict,
):

    if monto_por_persona <= 0:
        raise HTTPException(
            status_code=400,
            detail="El monto_por_persona debe ser mayor a 0"
        )

    condiciones, params = _build_filtros_where(db, filtros)

    where_sql = " AND ".join(condiciones) if condiciones else "1=1"

    sql = f"""
    SELECT
        COUNT(*) AS total
    FROM expediente_electronico e

    JOIN cuenta_bancaria_expediente cb
        ON cb.expediente_id = e.id

    LEFT JOIN (
        SELECT
            expediente_id,
            COUNT(*) AS total_pagos
        FROM expediente_cuenta_corriente
        WHERE tipo_movimiento = 'PAGO'
        GROUP BY expediente_id
    ) pagos
        ON pagos.expediente_id = e.id

    LEFT JOIN cat_estado_flujo_expediente cefe
        ON cefe.id = e.estado_flujo_id

    WHERE e.estado_flujo_id = 7
      AND {where_sql}
    """

    print("\n=========== SQL ===========")
    print(sql)
    print("PARAMS:", params)
    print("===========================\n")

    row = db.execute(text(sql), params).first()

    beneficiarios_encontrados = int(getattr(row, "total", 0) or 0)

    # cuantos permite el presupuesto
    beneficiarios_posibles = int(presupuesto_total // monto_por_persona)

    # cuantos realmente se pagarán
    beneficiarios_final = min(
        beneficiarios_encontrados,
        beneficiarios_posibles
    )

    monto_estimado = beneficiarios_final * monto_por_persona

    return {
        "beneficiarios_encontrados": beneficiarios_encontrados,
        "beneficiarios_posibles": beneficiarios_posibles,
        "monto_estimado": monto_estimado,
    }

def preview_detalle_lote_pago(
    db: Session,
    *,
    monto_por_persona: float,
    presupuesto_total: float,
    filtros: dict,
    page: int,
    limit: int,
):

    if monto_por_persona <= 0:
        raise HTTPException(
            status_code=400,
            detail="El monto_por_persona debe ser mayor a 0"
        )

    condiciones, params = _build_filtros_where(db, filtros)

    where_sql = " AND ".join(condiciones) if condiciones else "1=1"

    offset = (page - 1) * limit

    # ------------------------------
    # total de beneficiarios
    # ------------------------------

    sql_total = f"""
    SELECT
        COUNT(*) AS total
    FROM expediente_electronico e

    JOIN cuenta_bancaria_expediente cb
        ON cb.expediente_id = e.id

    LEFT JOIN (
        SELECT
            expediente_id,
            COUNT(*) AS total_pagos
        FROM expediente_cuenta_corriente
        WHERE tipo_movimiento = 'PAGO'
        GROUP BY expediente_id
    ) pagos
        ON pagos.expediente_id = e.id

    WHERE e.estado_flujo_id = 7
      AND {where_sql}
    """

    row = db.execute(text(sql_total), params).first()
    encontrados = int(getattr(row, "total", 0) or 0)

    beneficiarios_posibles = int(presupuesto_total // monto_por_persona)

    beneficiarios_final = min(
        encontrados,
        beneficiarios_posibles
    )

    # ------------------------------
    # query detalle
    # ------------------------------

    sql = f"""
    SELECT
        e.id AS expediente_id,
        e.nombre_beneficiario,
        e.cui_beneficiario,
        cb.numero_cuenta,
        cb.banco_codigo,
        d.nombre AS departamento,
        m.nombre AS municipio,
        COALESCE(pagos.total_pagos,0) AS total_pagos

    FROM expediente_electronico e

    JOIN cuenta_bancaria_expediente cb
        ON cb.expediente_id = e.id

    LEFT JOIN (
        SELECT
            expediente_id,
            COUNT(*) AS total_pagos
        FROM expediente_cuenta_corriente
        WHERE tipo_movimiento = 'PAGO'
        GROUP BY expediente_id
    ) pagos
        ON pagos.expediente_id = e.id

    LEFT JOIN cat_departamento d
        ON d.id = e.departamento_id

    LEFT JOIN cat_municipio m
        ON m.id = e.municipio_id

    WHERE e.estado_flujo_id = 7
      AND {where_sql}

    ORDER BY e.created_at DESC

    LIMIT :limit
    OFFSET :offset
    """

    params["limit"] = limit
    params["offset"] = offset

    rows = db.execute(text(sql), params).fetchall()

    data = []

    for r in rows:
        data.append({
            "id": r.expediente_id,
            "nombre_beneficiario": r.nombre_beneficiario,
            "cui_beneficiario": r.cui_beneficiario,
            "numero_cuenta": r.numero_cuenta,
            "banco_codigo": r.banco_codigo,
            "departamento": r.departamento,
            "municipio": r.municipio,
            "total_pagos": r.total_pagos,
        })

    return {
        "data": data,
        "page": page,
        "limit": limit,
        "total": beneficiarios_final
    }

# -------------------------------------------------
# BANDEJA: expedientes con cuenta bancaria
# -------------------------------------------------
def bandeja_expedientes_con_cuenta(
    db: Session,
    *,
    estado_flujo_codigo: str = "CUENTA_BANCARIA_CREADA",
    texto: Optional[str] = None,
    departamento_id: Optional[int] = None,
    municipio_id: Optional[int] = None,
    page: int = 1,
    limit: int = 20,
):
    texto = (texto or "").strip()

    q = (
        db.query(
            ExpedienteElectronico.id,
            ExpedienteElectronico.created_at,
            ExpedienteElectronico.nombre_beneficiario,
            ExpedienteElectronico.cui_beneficiario,
            ExpedienteElectronico.estado_expediente,
            ExpedienteElectronico.bpm_status,
            ExpedienteElectronico.bpm_current_task_name,
            CatEstadoFlujoExpediente.codigo.label("estado_flujo_codigo"),
            CatEstadoFlujoExpediente.nombre.label("estado_flujo_nombre"),
            CatDepartamento.nombre.label("departamento"),
            CatMunicipio.nombre.label("municipio"),
            CuentaBancariaExpediente.banco_codigo.label("banco_codigo"),
            CuentaBancariaExpediente.numero_cuenta.label("numero_cuenta"),
        )
        .join(CuentaBancariaExpediente, CuentaBancariaExpediente.expediente_id == ExpedienteElectronico.id)
        .outerjoin(CatEstadoFlujoExpediente, CatEstadoFlujoExpediente.id == ExpedienteElectronico.estado_flujo_id)
        .outerjoin(CatDepartamento, CatDepartamento.id == ExpedienteElectronico.departamento_id)
        .outerjoin(CatMunicipio, CatMunicipio.id == ExpedienteElectronico.municipio_id)
    )

    filters = []

    if estado_flujo_codigo:
        filters.append(CatEstadoFlujoExpediente.codigo == estado_flujo_codigo)

    if texto:
        filters.append(
            or_(
                ExpedienteElectronico.nombre_beneficiario.ilike(f"%{texto}%"),
                ExpedienteElectronico.cui_beneficiario.like(f"{texto}%"),
            )
        )
    if departamento_id:
        filters.append(ExpedienteElectronico.departamento_id == departamento_id)
    if municipio_id:
        filters.append(ExpedienteElectronico.municipio_id == municipio_id)

    total = q.with_entities(func.count(ExpedienteElectronico.id)).filter(*filters).scalar() or 0
    offset = (page - 1) * limit

    rows = (
        q.filter(*filters)
        .order_by(ExpedienteElectronico.created_at.desc())
        .offset(offset)
        .limit(limit)
        .all()
    )

    data = []
    for r in rows:
        data.append({
            "id": r.id,
            "created_at": r.created_at,
            "nombre_beneficiario": r.nombre_beneficiario,
            "cui_beneficiario": r.cui_beneficiario,
            "estado_expediente": r.estado_expediente,
            "bpm_status": r.bpm_status,
            "bpm_current_task_name": r.bpm_current_task_name,
            "estado_flujo_codigo": getattr(r, "estado_flujo_codigo", None),
            "estado_flujo_nombre": getattr(r, "estado_flujo_nombre", None),
            "departamento": r.departamento,
            "municipio": r.municipio,
            "banco_codigo": getattr(r, "banco_codigo", None),
            "numero_cuenta": getattr(r, "numero_cuenta", None),
        })

    return {"data": data, "page": page, "limit": limit, "total": total}


# -------------------------------------------------
# CREAR LOTE + ITEMS (snapshot mínimo)
# -------------------------------------------------
def crear_lote_pago(
    db: Session,
    *,
    expediente_ids: List[int],
    anio_fiscal: int,
    mes_fiscal: int,
    monto_por_persona: float,
    tope_anual_persona: float,
    creado_por: Optional[str] = None,
    observacion: Optional[str] = None,
) -> Tuple[int, int]:

    rows = (
        db.query(
            ExpedienteElectronico.id,
            ExpedienteElectronico.cui_beneficiario,
            ExpedienteElectronico.nombre_beneficiario,
            CatDepartamento.nombre.label("departamento"),
            CatMunicipio.nombre.label("municipio"),
            CuentaBancariaExpediente.banco_codigo.label("banco_codigo"),
            CuentaBancariaExpediente.numero_cuenta.label("numero_cuenta"),
        )
        .join(CuentaBancariaExpediente, CuentaBancariaExpediente.expediente_id == ExpedienteElectronico.id)
        .outerjoin(CatDepartamento, CatDepartamento.id == ExpedienteElectronico.departamento_id)
        .outerjoin(CatMunicipio, CatMunicipio.id == ExpedienteElectronico.municipio_id)
        .filter(ExpedienteElectronico.id.in_(expediente_ids))
        .all()
    )

    found_ids = {r.id for r in rows}
    missing = [i for i in expediente_ids if i not in found_ids]
    if missing:
        raise HTTPException(status_code=404, detail=f"Expedientes no encontrados o sin cuenta: {missing[:20]}")

    lote = LotePago(
        anio_fiscal=anio_fiscal,
        mes_fiscal=mes_fiscal,
        monto_por_persona=monto_por_persona,
        tope_anual_persona=tope_anual_persona,
        banco_codigo="BANRURAL",
        estado="CREADO",
        creado_por=creado_por,
        observacion=observacion,
    )
    db.add(lote)
    db.flush()

    # ✅ Crear items
    for r in rows:
        db.add(DetallePago(
            lote_id=lote.id,
            expediente_id=r.id,
            anio_fiscal=anio_fiscal,
            mes_fiscal=mes_fiscal,
            estado="PENDIENTE",
            monto_asignado=monto_por_persona,
            cui_beneficiario=r.cui_beneficiario,
            nombre_beneficiario=r.nombre_beneficiario,
            banco_codigo=r.banco_codigo or "BANRURAL",
            numero_cuenta=r.numero_cuenta,
        ))

        # ✅ TRACKING IMPORTANTE: intento de pago (por expediente)
        TrackingEventoService._registrar(
            db,
            expediente_id=int(r.id),
            titulo="Intento de pago",
            origen=TrackingEventoService.ORIGEN_PAGOS,
            tipo_evento=TrackingEventoService.PAGO_INTENTO,
            usuario=creado_por,
            observacion=f"Lote #{lote.id} | Periodo {anio_fiscal}-{mes_fiscal:02d} | Monto Q{float(monto_por_persona):,.2f}",
            commit=False,
        )

    db.commit()
    return lote.id, len(rows)

# -------------------------------------------------
# LISTAR LOTES (paginado + filtros)
# -------------------------------------------------
def listar_lotes_pago(
    db: Session,
    *,
    anio: Optional[int] = None,
    mes: Optional[int] = None,
    estado: Optional[str] = None,
    texto: Optional[str] = None,
    page: int = 1,
    limit: int = 20,
):
    texto = (texto or "").strip()
    offset = (page - 1) * limit

    q = db.query(LotePago)

    q = q.filter(LotePago.estado != "ELIMINADO")

    if anio:
        q = q.filter(LotePago.anio_fiscal == anio)
    if mes:
        q = q.filter(LotePago.mes_fiscal == mes)
    if estado:
        q = q.filter(LotePago.estado == estado)

    if texto:
        q = q.filter(
            or_(
                func.cast(LotePago.id, String).ilike(f"%{texto}%"),
                LotePago.estado.ilike(f"%{texto}%"),
                (LotePago.creado_por.ilike(f"%{texto}%") if LotePago.creado_por is not None else False),
                (LotePago.observacion.ilike(f"%{texto}%") if LotePago.observacion is not None else False),
            )
        )

    total = q.with_entities(func.count(LotePago.id)).scalar() or 0
    lotes = q.order_by(LotePago.creado_en.desc()).offset(offset).limit(limit).all()

    lote_ids = [l.id for l in lotes]
    counts_map = {lid: {"total_items": 0, "pagados": 0, "rechazados": 0} for lid in lote_ids}

    if lote_ids:
        agg = (
            db.query(
                DetallePago.lote_id.label("lote_id"),
                func.count(DetallePago.id).label("total_items"),
                func.coalesce(func.sum(case((DetallePago.estado == "PAGADO", 1), else_=0)), 0).label("pagados"),
                func.coalesce(func.sum(case((DetallePago.estado == "RECHAZADO", 1), else_=0)), 0).label("rechazados"),
            )
            .filter(DetallePago.lote_id.in_(lote_ids))
            .group_by(DetallePago.lote_id)
            .all()
        )
        for r in agg:
            counts_map[r.lote_id] = {
                "total_items": int(r.total_items or 0),
                "pagados": int(r.pagados or 0),
                "rechazados": int(r.rechazados or 0),
            }

    data = []
    for l in lotes:
        c = counts_map.get(l.id, {"total_items": 0, "pagados": 0, "rechazados": 0})
        data.append({
            "id": l.id,
            "anio_fiscal": l.anio_fiscal,
            "mes_fiscal": l.mes_fiscal,
            "banco_codigo": l.banco_codigo,
            "estado": l.estado,
            "creado_por": l.creado_por,
            "creado_en": l.creado_en,
            "procesado_en": l.procesado_en,
            "observacion": l.observacion,
            **c,
        })

    return {"data": data, "page": page, "limit": limit, "total": total}


# -------------------------------------------------
# OBTENER DETALLE DE LOTE
# -------------------------------------------------
def obtener_lote_pago(db: Session, *, lote_id: int) -> dict:

    lote = db.query(LotePago).filter(LotePago.id == lote_id).first()

    if not lote:
        raise HTTPException(status_code=404, detail="Planilla no encontrada")

    agg = (
        db.query(
            func.count(DetallePago.id).label("total_items"),
            func.coalesce(
                func.sum(case((DetallePago.estado == "PAGADO", 1), else_=0)), 0
            ).label("pagados"),
            func.coalesce(
                func.sum(case((DetallePago.estado == "RECHAZADO", 1), else_=0)), 0
            ).label("rechazados"),
        )
        .filter(DetallePago.lote_id == lote_id)
        .first()
    )

    presupuesto_total = (
        float(lote.presupuesto_total)
        if getattr(lote, "presupuesto_total", None) is not None
        else None
    )

    monto_usado = (
        float(lote.monto_usado)
        if getattr(lote, "monto_usado", None) is not None
        else None
    )

    monto_no_usado = None
    if presupuesto_total is not None and monto_usado is not None:
        monto_no_usado = presupuesto_total - monto_usado

    filtros_resueltos = resolver_filtros(db, lote.filtros_json)

    return {

        "id": lote.id,

        "anio_fiscal": lote.anio_fiscal,
        "mes_fiscal": lote.mes_fiscal,

        "banco_codigo": lote.banco_codigo,
        "estado": lote.estado,

        "creado_por": lote.creado_por,
        "creado_en": lote.creado_en,
        "procesado_en": lote.procesado_en,

        "observacion": lote.observacion,

        "monto_por_persona": float(lote.monto_por_persona),
        "tope_anual_persona": float(lote.tope_anual_persona),

        "presupuesto_total": presupuesto_total,
        "monto_usado": monto_usado,
        "monto_no_usado": monto_no_usado,

        "beneficiarios_encontrados": getattr(lote, "beneficiarios_encontrados", None),

        "filtros_json": lote.filtros_json,
        "filtros_resueltos": filtros_resueltos,

        "total_items": int(getattr(agg, "total_items", 0) or 0),
        "pagados": int(getattr(agg, "pagados", 0) or 0),
        "rechazados": int(getattr(agg, "rechazados", 0) or 0),

    }

# =====================================================
# 🔹 Resolver filtros (IDs → nombres)
# =====================================================
def resolver_filtros(db: Session, filtros_json: dict | None) -> dict | None:

    if not filtros_json:
        return None

    numero_pago = filtros_json.get("numero_pago")
    ubicaciones = filtros_json.get("ubicaciones")

    # ===============================
    # 🔹 Resolver numero_pago
    # ===============================
    numero_pago_resuelto = numero_pago

    # ===============================
    # 🔹 Resolver ubicaciones
    # ===============================
    ubicaciones_resueltas = []

    if ubicaciones:

        # 🔹 Obtener todos los catálogos de una vez
        departamentos = {
            d.id: d.nombre
            for d in db.query(CatDepartamento).all()
        }

        municipios = db.query(CatMunicipio).all()

        municipios_map = {}
        for m in municipios:
            municipios_map.setdefault(m.departamento_id, []).append(m)

        # 🔹 Mapear
        for u in ubicaciones:

            depto_id = u.get("departamento_id")
            municipios_ids = u.get("municipios_ids")

            depto_nombre = departamentos.get(depto_id, f"ID {depto_id}")

            if municipios_ids:
                nombres_municipios = [
                    m.nombre
                    for m in municipios_map.get(depto_id, [])
                    if m.id in municipios_ids
                ]
            else:
                nombres_municipios = None  # Todos

            ubicaciones_resueltas.append({
                "departamento": depto_nombre,
                "municipios": nombres_municipios
            })

    return {
        "numero_pago": numero_pago_resuelto,
        "ubicaciones": ubicaciones_resueltas if ubicaciones else None
    }

# -------------------------------------------------
# LISTAR ITEMS DE LOTE
# -------------------------------------------------
def listar_items_lote_pago(db: Session, *, lote_id: int, page: int = 1, limit: int = 50) -> dict:
    lote = db.query(LotePago).filter(LotePago.id == lote_id).first()  # CAMBIO
    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    base_q = (
        db.query(
            DetallePago,
            ExpedienteElectronico,
            CatDepartamento,  # CAMBIO
            CatMunicipio,  # CAMBIO
        )
        .outerjoin(
            ExpedienteElectronico,
            ExpedienteElectronico.id == DetallePago.expediente_id,
        )
        .outerjoin(  # CAMBIO
            CatDepartamento,
            CatDepartamento.id == ExpedienteElectronico.departamento_id,
        )
        .outerjoin(  # CAMBIO
            CatMunicipio,
            CatMunicipio.id == ExpedienteElectronico.municipio_id,
        )
        .filter(DetallePago.lote_id == lote_id)
    )

    total = (
        db.query(func.count(DetallePago.id))
        .filter(DetallePago.lote_id == lote_id)
        .scalar()
        or 0
    )  # CAMBIO

    offset = (page - 1) * limit

    rows = (
        base_q.order_by(DetallePago.id.asc())
        .offset(offset)
        .limit(limit)
        .all()
    )

    data = []
    for it, exp, dep, mun in rows:  # CAMBIO
        data.append({
            "id": it.id,
            "lote_id": it.lote_id,
            "expediente_id": it.expediente_id,

            "departamento": dep.nombre if dep else None,  # CAMBIO
            "municipio": mun.nombre if mun else None,  # CAMBIO

            "anio_fiscal": it.anio_fiscal,
            "mes_fiscal": it.mes_fiscal,
            "estado": it.estado,
            "monto_asignado": float(it.monto_asignado),

            "numero_pago": getattr(lote, "numero_pago", None),  # CAMBIO

            "acumulado_pagado_antes": float(it.acumulado_pagado_antes) if it.acumulado_pagado_antes is not None else None,
            "excede_tope": bool(it.excede_tope),
            "cui_beneficiario": it.cui_beneficiario,
            "nombre_beneficiario": it.nombre_beneficiario,
            "banco_codigo": it.banco_codigo,

            "titular": exp.titular_nombre if exp else None,  # CAMBIO
            "numero_cuenta": it.numero_cuenta,

            "motivo_rechazo": it.motivo_rechazo,
            "referencia_externa": it.referencia_externa,
            "procesado_en": it.procesado_en,
        })

    return {"data": data, "page": page, "limit": limit, "total": total}

def generar_excel_lote_pago(db: Session, *, lote_id: int) -> BytesIO:

    exists = db.query(LotePago.id).filter(LotePago.id == lote_id).first()
    if not exists:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    rows = (
        db.query(DetallePago)
        .filter(DetallePago.lote_id == lote_id)
        .order_by(DetallePago.id.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Planilla"

    headers = [
        "EXPEDIENTE",
        "CUI",
        "NOMBRE",
        "CUENTA",
        "BANCO",
        "MONTO",
        "ESTADO",
        "ACUMULADO ANTES",
        "EXCEDE TOPE",
        "MOTIVO RECHAZO",
        "REFERENCIA",
        "FECHA PROCESO",
    ]

    ws.append(headers)

    # 🔹 estilo header
    for col in ws[1]:
        col.font = Font(bold=True)

    for it in rows:
        ws.append([
            it.expediente_id,
            it.cui_beneficiario,
            it.nombre_beneficiario,
            it.numero_cuenta,
            it.banco_codigo,
            float(it.monto_asignado),
            it.estado,
            float(it.acumulado_pagado_antes) if it.acumulado_pagado_antes else None,
            "SI" if it.excede_tope else "NO",
            it.motivo_rechazo,
            it.referencia_externa,
            it.procesado_en.strftime("%d/%m/%Y %H:%M") if it.procesado_en else None,
        ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return stream

def generar_excel_beneficiarios_por_ids(
    db: Session,
    *,
    beneficiario_ids: list[int],
) -> BytesIO:

    # =========================
    # VALIDACIONES
    # =========================
    if not beneficiario_ids:
        raise HTTPException(
            status_code=400,
            detail="Debe enviar beneficiario_ids"
        )

    beneficiario_ids = list({int(x) for x in beneficiario_ids if x is not None})

    if not beneficiario_ids:
        raise HTTPException(
            status_code=400,
            detail="No hay beneficiarios válidos para exportar"
        )

    # =========================
    # CONSULTA
    # =========================
    rows = db.execute(
        text("""
            SELECT
                e.id AS expediente_id,
                e.cui_beneficiario,
                e.nombre_beneficiario,
                e.titular_nombre,
                e.titular_dpi,
                cb.numero_cuenta,
                d.nombre AS departamento,
                m.nombre AS municipio,
                e.created_at
            FROM expediente_electronico e
            JOIN cuenta_bancaria_expediente cb
                ON cb.expediente_id = e.id
            LEFT JOIN cat_departamento d
                ON d.id = e.departamento_id
            LEFT JOIN cat_municipio m
                ON m.id = e.municipio_id
            WHERE e.id = ANY(:ids)
            ORDER BY e.created_at DESC
        """),
        {"ids": beneficiario_ids}
    ).fetchall()

    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No se encontraron beneficiarios"
        )

    # =========================
    # EXCEL
    # =========================
    wb = Workbook()
    ws = wb.active
    ws.title = "Beneficiarios"

    headers = [
        "EXPEDIENTE",
        "CUI",
        "NOMBRE",
        "TITULAR",
        "DPI TITULAR",
        "CUENTA",
        "DEPARTAMENTO",
        "MUNICIPIO",
        "FECHA CREACIÓN",
    ]

    ws.append(headers)

    for col in ws[1]:
        col.font = XLFont(bold=True)

    for r in rows:
        ws.append([
            r.expediente_id,
            r.cui_beneficiario,
            r.nombre_beneficiario,
            r.titular_nombre,
            r.titular_dpi,
            r.numero_cuenta,
            r.departamento,
            r.municipio,
            r.created_at.strftime("%d/%m/%Y %H:%M") if r.created_at else None,
        ])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return stream

# -------------------------------------------------
# PROCESAR LOTE (SIMULADO) + VALIDACIÓN TOPE ANUAL
# -------------------------------------------------
def procesar_lote_pago_simulado(db: Session, *, lote_id: int) -> dict:
    lote = db.query(LotePago).filter(LotePago.id == lote_id).first()
    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    if (lote.estado or "").upper() == "PROCESADO":
        # Idempotente: devuelve resumen actual (no vuelve a registrar tracking)
        agg = (
            db.query(
                func.count(DetallePago.id).label("total_items"),
                func.coalesce(func.sum(case((DetallePago.estado == "PAGADO", 1), else_=0)), 0).label("pagados"),
                func.coalesce(func.sum(case((DetallePago.estado == "RECHAZADO", 1), else_=0)), 0).label("rechazados"),
            )
            .filter(DetallePago.lote_id == lote_id)
            .first()
        )
        return {
            "lote_id": lote_id,
            "total_items": int(getattr(agg, "total_items", 0) or 0),
            "pagados": int(getattr(agg, "pagados", 0) or 0),
            "rechazados": int(getattr(agg, "rechazados", 0) or 0),
            "procesado_en": lote.procesado_en,
        }

    now = datetime.utcnow()

    items = (
        db.query(DetallePago)
        .filter(DetallePago.lote_id == lote_id)
        .order_by(DetallePago.id.asc())
        .all()
    )

    total_items = len(items)
    pagados = 0
    rechazados = 0

    for it in items:
        if (it.estado or "").upper() not in ("PENDIENTE", ""):
            continue

        try:
            # acumulado anual pagado:
            # 1) histórico real en cuenta corriente del expediente
            acumulado_cc = (
                db.query(func.coalesce(func.sum(ExpedienteCuentaCorriente.monto), 0))
                .filter(
                    ExpedienteCuentaCorriente.expediente_id == it.expediente_id,
                    ExpedienteCuentaCorriente.tipo_movimiento == "PAGO",
                )
                .scalar()
                or 0
            )

            # 2) por compatibilidad, también se conserva el cálculo previo sobre detalle_pago pagado del mismo año
            acumulado_detalle = (
                db.query(func.coalesce(func.sum(DetallePago.monto_asignado), 0))
                .filter(
                    DetallePago.expediente_id == it.expediente_id,
                    DetallePago.anio_fiscal == lote.anio_fiscal,
                    DetallePago.estado == "PAGADO",
                )
                .scalar()
                or 0
            )

            acumulado = max(float(acumulado_cc), float(acumulado_detalle))

            it.acumulado_pagado_antes = acumulado
            it.actualizado_en = now

            monto = float(it.monto_asignado)
            tope = float(lote.tope_anual_persona)

            if (acumulado + monto) > tope:
                it.estado = "RECHAZADO"
                it.excede_tope = True
                it.motivo_rechazo = "EXCEDE_TOPE_ANUAL"
                it.response_status_code = 409
                it.response_payload = {
                    "ok": False,
                    "reason": "EXCEDE_TOPE_ANUAL",
                    "acumulado": float(acumulado),
                    "monto": monto,
                    "tope": tope,
                }
                it.procesado_en = now
                rechazados += 1

                # ✅ TRACKING IMPORTANTE: excede tope (por expediente)
                TrackingEventoService._registrar(
                    db,
                    expediente_id=int(it.expediente_id),
                    titulo="Pago rechazado por exceder tope anual",
                    origen=TrackingEventoService.ORIGEN_PAGOS,
                    tipo_evento=TrackingEventoService.PAGO_EXCEDE_TOPE,
                    usuario=None,
                    observacion=f"Lote #{lote.id} | Periodo {lote.anio_fiscal}-{lote.mes_fiscal:02d} | Acumulado Q{float(acumulado):,.2f} | Monto Q{monto:,.2f} | Tope Q{tope:,.2f}",
                    commit=False,
                )
                continue

            # ✅ Simulación de “consumo”
            ref = f"SIM-{uuid4().hex[:12].upper()}"
            it.estado = "PAGADO"
            it.excede_tope = False
            it.motivo_rechazo = None
            it.referencia_externa = ref
            it.proveedor_servicio = it.proveedor_servicio or lote.proveedor_servicio or "SIMULADO"
            it.response_status_code = 200
            it.response_payload = {"ok": True, "referencia": ref}
            it.procesado_en = now
            pagados += 1

            # ✅ REGISTRO EN CUENTA CORRIENTE
            db.add(
                ExpedienteCuentaCorriente(
                    expediente_id=it.expediente_id,
                    tipo_movimiento="PAGO",
                    referencia_tipo="LOTE_PAGO",
                    referencia_id=lote.id,
                    monto=it.monto_asignado,
                    descripcion=f"Pago lote #{lote.id} periodo {lote.anio_fiscal}-{lote.mes_fiscal:02d}",
                    created_by=None,
                )
            )

            # ✅ TRACKING IMPORTANTE: pago aprobado (por expediente)
            TrackingEventoService._registrar(
                db,
                expediente_id=int(it.expediente_id),
                titulo="Pago aprobado",
                origen=TrackingEventoService.ORIGEN_PAGOS,
                tipo_evento=TrackingEventoService.PAGO_APROBADO,
                usuario=None,
                observacion=f"Lote #{lote.id} | Periodo {lote.anio_fiscal}-{lote.mes_fiscal:02d} | Monto Q{monto:,.2f} | Ref {ref}",
                commit=False,
            )

            # ---------------------------------------------------------
            # (FUTURO) Consumo real del servicio externo (DEJAR COMENTADO)
            # ---------------------------------------------------------
            # payload = {...}
            # resp = external_client.post(..., json=payload)
            # it.request_payload = payload
            # it.response_payload = resp.json()
            # it.response_status_code = resp.status_code
            # if resp.ok: PAGADO else RECHAZADO
            # ---------------------------------------------------------

        except Exception as e:
            # ✅ TRACKING IMPORTANTE: error técnico (por expediente)
            TrackingEventoService._registrar(
                db,
                expediente_id=int(it.expediente_id),
                titulo="Error técnico en proceso de pago",
                origen=TrackingEventoService.ORIGEN_PAGOS,
                tipo_evento=TrackingEventoService.PAGO_ERROR,
                usuario=None,
                observacion=str(e)[:500],
                commit=False,
            )

            # Estado de item (si quieres mantenerlo pendiente o marcarlo rechazado)
            it.estado = "RECHAZADO"
            it.excede_tope = False
            it.motivo_rechazo = "ERROR_TECNICO"
            it.response_status_code = 500
            it.response_payload = {"ok": False, "reason": "ERROR_TECNICO"}
            it.procesado_en = now
            rechazados += 1

    lote.estado = "PROCESADO"
    lote.procesado_en = now
    lote.total_items = total_items
    lote.pagados = pagados
    lote.rechazados = rechazados

    db.commit()

    return {
        "lote_id": lote_id,
        "total_items": total_items,
        "pagados": pagados,
        "rechazados": rechazados,
        "procesado_en": now,
    }


# -------------------------------------------------
# EXCEL EXPORT (NO solicitud, solo export evidencia)
# -------------------------------------------------
def generar_excel_lote_pago_export_bytes(db: Session, *, lote_id: int) -> bytes:
    lote = db.query(LotePago).filter(LotePago.id == lote_id).first()
    if not lote:
        raise HTTPException(status_code=404, detail="Lote no encontrado")

    items = (
        db.query(DetallePago)
        .filter(DetallePago.lote_id == lote_id)
        .order_by(DetallePago.id.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "PAGOS_EXPORT"

    headers = [
        "lote_id",
        "anio_fiscal",
        "mes_fiscal",
        "expediente_id",
        "cui_beneficiario",
        "nombre_beneficiario",
        "banco_codigo",
        "numero_cuenta",
        "monto_asignado",
        "acumulado_pagado_antes",
        "tope_anual",
        "estado",
        "motivo_rechazo",
        "referencia_externa",
        "procesado_en",
        "creado_en",
    ]
    ws.append(headers)

    for it in items:
        ws.append([
            lote.id,
            lote.anio_fiscal,
            lote.mes_fiscal,
            it.expediente_id,
            it.cui_beneficiario,
            it.nombre_beneficiario,
            it.banco_codigo,
            it.numero_cuenta,
            float(it.monto_asignado),
            float(it.acumulado_pagado_antes) if it.acumulado_pagado_antes is not None else None,
            float(lote.tope_anual_persona),
            it.estado,
            it.motivo_rechazo,
            it.referencia_externa,
            it.procesado_en.isoformat() if it.procesado_en else None,
            it.creado_en.isoformat() if it.creado_en else None,
        ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()

def crear_lote_pago_por_filtros(
    db: Session,
    *,
    anio_fiscal: int,
    mes_fiscal: int,
    monto_por_persona: float,
    tope_anual_persona: float,
    presupuesto_total: float,

    # CAMBIO 🔥
    filtros: list,
    beneficiario_ids_unicos: list[int],

    creado_por: Optional[str] = None,
    observacion: Optional[str] = None,
) -> Tuple[int, int]:

    # =========================
    # VALIDACIONES
    # =========================

    if monto_por_persona <= 0:
        raise HTTPException(
            status_code=400,
            detail="monto_por_persona debe ser mayor a 0"
        )

    if not beneficiario_ids_unicos:
        raise HTTPException(
            status_code=400,
            detail="Debe enviar beneficiarios"
        )

    max_beneficiarios = int(presupuesto_total // monto_por_persona)

    if max_beneficiarios <= 0:
        raise HTTPException(
            status_code=400,
            detail="El presupuesto no alcanza para ningún beneficiario"
        )

    total_actual = len(beneficiario_ids_unicos)

    # =========================
    # VALIDACIÓN PRESUPUESTO
    # =========================
    if total_actual > max_beneficiarios:
        raise HTTPException(
            status_code=409,
            detail=(
                f"El presupuesto permite {max_beneficiarios} beneficiarios, "
                f"pero se enviaron {total_actual}. "
                "Por favor, revise los filtros antes de continuar."
            )
        )

    # =========================
    # 🔥 OBTENER DATOS REALES (SOLO IDs)
    # =========================
    rows = db.execute(
        text("""
            SELECT
                e.id,
                e.nombre_beneficiario,
                e.cui_beneficiario,
                cb.numero_cuenta,
                cb.banco_codigo
            FROM expediente_electronico e
            JOIN cuenta_bancaria_expediente cb
                ON cb.expediente_id = e.id
            WHERE e.id = ANY(:ids)
        """),
        {"ids": beneficiario_ids_unicos}
    ).fetchall()

    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No se encontraron beneficiarios"
        )

    # =========================
    # CREAR LOTE
    # =========================
    lote = LotePago(
        anio_fiscal=anio_fiscal,
        mes_fiscal=mes_fiscal,
        monto_por_persona=monto_por_persona,
        tope_anual_persona=tope_anual_persona,
        presupuesto_total=presupuesto_total,

        # CAMBIO 🔥
        monto_usado=len(rows) * monto_por_persona,
        beneficiarios_encontrados=len(rows),

        # CAMBIO 🔥 (guardamos TODO)
        filtros_json={
            "filtros": filtros,
            "beneficiario_ids_unicos": beneficiario_ids_unicos,
        },

        banco_codigo="BANRURAL",
        estado="CREADO",
        creado_por=creado_por,
        observacion=observacion,
    )

    db.add(lote)
    db.flush()

    # =========================
    # CREAR ITEMS
    # =========================
    for r in rows:

        db.add(
            DetallePago(
                lote_id=lote.id,
                expediente_id=r.id,
                anio_fiscal=anio_fiscal,
                mes_fiscal=mes_fiscal,
                estado="PENDIENTE",
                monto_asignado=monto_por_persona,
                cui_beneficiario=r.cui_beneficiario,
                nombre_beneficiario=r.nombre_beneficiario,
                banco_codigo=r.banco_codigo or "BANRURAL",
                numero_cuenta=r.numero_cuenta,
            )
        )

        TrackingEventoService._registrar(
            db,
            expediente_id=int(r.id),
            titulo="Intento de pago",
            origen=TrackingEventoService.ORIGEN_PAGOS,
            tipo_evento=TrackingEventoService.PAGO_INTENTO,
            usuario=creado_por,
            observacion=(
                f"Lote #{lote.id} | Periodo {anio_fiscal}-{mes_fiscal:02d} "
                f"| Monto Q{float(monto_por_persona):,.2f}"
            ),
            commit=False,
        )

    db.commit()

    return lote.id, len(rows)

# -------------------------------------------------
# OBTENER BENEFICIARIOS CON FILTROS
# -------------------------------------------------

def obtener_beneficiarios_para_pago(
    db: Session,
    *,
    numero_pago: int,
    ubicaciones: List[Dict[str, Any]] | None = None,
    limite: int | None = None,
) -> List[Dict[str, Any]]:

    # ------------------------------
    # Validación
    # ------------------------------
    if numero_pago < 0 or numero_pago > 11:
        raise ValueError("numero_pago debe estar entre 0 y 11")

    # ------------------------------
    # Filtro de ubicaciones
    # ------------------------------
    filtros_ubicacion = []
    params = {
        "numero_pago": numero_pago
    }

    # 🚫 Regla: pago 0 NO usa ubicación
    if numero_pago != 0 and ubicaciones:

        condiciones = []

        for i, ub in enumerate(ubicaciones):
            dep_key = f"dep_{i}"
            params[dep_key] = ub["departamento_id"]

            municipios = ub.get("municipios") or []

            # 🔹 Departamento completo
            if not municipios:
                condiciones.append(f"(e.departamento_id = :{dep_key})")

            # 🔹 Departamento + municipios
            else:
                mun_keys = []
                for j, m in enumerate(municipios):
                    key = f"mun_{i}_{j}"
                    params[key] = m
                    mun_keys.append(f":{key}")

                condiciones.append(
                    f"(e.departamento_id = :{dep_key} AND e.municipio_id IN ({','.join(mun_keys)}))"
                )

        if condiciones:
            filtros_ubicacion.append("(" + " OR ".join(condiciones) + ")")

    # ------------------------------
    # WHERE FINAL
    # ------------------------------
    where_extra = ""
    if filtros_ubicacion:
        where_extra = " AND " + " AND ".join(filtros_ubicacion)

    # ------------------------------
    # SQL BASE
    # ------------------------------
    sql = f"""
    SELECT
        e.id AS expediente_id,
        e.nombre_beneficiario,
        e.cui_beneficiario,
        cb.numero_cuenta,
        cb.banco_codigo,
        d.nombre AS departamento,
        m.nombre AS municipio,
        COALESCE(pagos.total_pagos, 0) AS total_pagos,
        e.created_at

    FROM expediente_electronico e

    JOIN cuenta_bancaria_expediente cb
        ON cb.expediente_id = e.id

    LEFT JOIN (
        SELECT
            expediente_id,
            COUNT(*) AS total_pagos
        FROM expediente_cuenta_corriente
        WHERE tipo_movimiento = 'PAGO'
        GROUP BY expediente_id
    ) pagos
        ON pagos.expediente_id = e.id

    LEFT JOIN cat_departamento d
        ON d.id = e.departamento_id

    LEFT JOIN cat_municipio m
        ON m.id = e.municipio_id

    WHERE e.estado_flujo_id = 7
      AND COALESCE(pagos.total_pagos, 0) = :numero_pago
      {where_extra}

    ORDER BY e.created_at DESC
    """

    # ------------------------------
    # Límite opcional (Excel / presupuesto)
    # ------------------------------
    if limite is not None:
        sql += "\nLIMIT :limite"
        params["limite"] = limite

    # ------------------------------
    # Ejecutar
    # ------------------------------
    rows = db.execute(text(sql), params).fetchall()

    # ------------------------------
    # Transformar
    # ------------------------------
    data = []

    for r in rows:
        data.append({
            "id": r.expediente_id,
            "nombre_beneficiario": r.nombre_beneficiario,
            "cui_beneficiario": r.cui_beneficiario,
            "numero_cuenta": r.numero_cuenta,
            "banco_codigo": r.banco_codigo,
            "departamento": r.departamento,
            "municipio": r.municipio,
            "total_pagos": r.total_pagos,
            "created_at": r.created_at,
        })

    return data

def obtener_totales_beneficiarios_pago(
    db: Session,
    *,
    numero_pago: int,
    ubicaciones: list | None,
):

    # ------------------------------
    # Validación
    # ------------------------------
    if numero_pago < 0 or numero_pago > 11:
        raise HTTPException(
            status_code=400,
            detail="numero_pago debe estar entre 0 y 11"
        )

    # ------------------------------
    # CORE (REUTILIZA TU FUNCIÓN)
    # ------------------------------
    rows = obtener_beneficiarios_para_pago(
        db,
        numero_pago=numero_pago,
        ubicaciones=ubicaciones,
        limite=None
    )

    if not rows:
        return {
            "beneficiarios_encontrados": 0,
            "beneficiarios_posibles": 0,
            "beneficiario_ids": [],  # CAMBIO
        }

    total = len(rows)

    beneficiario_ids = list({row["id"] for row in rows if row.get("id") is not None})  # CAMBIO

    return {
        "beneficiarios_encontrados": total,
        "beneficiarios_posibles": total,
        "beneficiario_ids": beneficiario_ids,  # CAMBIO
    }

def generar_excel_beneficiarios_pago(
    db: Session,
    *,
    numero_pago: int,
    ubicaciones: list | None,
):

    rows = obtener_beneficiarios_para_pago(
        db,
        numero_pago=numero_pago,
        ubicaciones=ubicaciones,
        limite=None
    )

    if not rows:
        raise HTTPException(
            status_code=404,
            detail="No se encontraron beneficiarios"
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "BENEFICIARIOS"

    # HEADERS
    headers = [
        "Expediente",
        "Nombre",
        "CUI",
        "Cuenta",
        "Banco",
        "Departamento",
        "Municipio",
        "Cantidad Pagos",
        "Fecha Creación",
    ]

    ws.append(headers)

    # DATA
    for r in rows:
        ws.append([
            r["id"],
            r["nombre_beneficiario"],
            r["cui_beneficiario"],
            r["numero_cuenta"],
            r["banco_codigo"],
            r["departamento"],
            r["municipio"],
            r["total_pagos"],
            str(r["created_at"]) if r["created_at"] else "",
        ])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=beneficiarios_pago_{numero_pago}.xlsx"
        }
    )

def generar_txt_resumen_lote(db: Session, *, lote_id: int):

    rows = (
        db.query(
            InfoGeneral.municipio_residencia_id.label("municipio_id"),
            func.count(DetallePago.id).label("total_pagos"),
            func.sum(DetallePago.monto_asignado).label("monto_total"),
        )
        .join(
            ExpedienteElectronico,
            ExpedienteElectronico.id == DetallePago.expediente_id
        )
        .join(
            InfoGeneral,
            InfoGeneral.expediente_id == ExpedienteElectronico.id
        )
        .filter(DetallePago.lote_id == lote_id)
        .group_by(InfoGeneral.municipio_residencia_id)
        .all()
    )

    if not rows:
        raise HTTPException(status_code=404, detail="No hay registros para este lote")

    lines = []

    # HEADER
    lines.append("GrupoPago,TipoPago,CodigoMunicipio,TotalPagos,MontoTotal,Programa")

    for r in rows:

        if r.municipio_id is None:
            raise HTTPException(
                status_code=400,
                detail="Existen beneficiarios sin municipio asignado"
            )

        line = ",".join([
            str(lote_id),                     
            "162078",                         
            str(r.municipio_id),              
            str(int(r.total_pagos)),          
            f"{float(r.monto_total):.2f}",    
            "2",                              
        ])

        lines.append(line)

    content = "\n".join(lines)

    buffer = io.BytesIO(content.encode("utf-8"))

    return StreamingResponse(
        buffer,
        media_type="text/plain",
        headers={
            "Content-Disposition": f"attachment; filename=resumen_lote_{lote_id}.csv"
        }
    )

def generar_txt_detalle_lote(db: Session, *, lote_id: int):

    items = (
        db.query(DetallePago)
        .filter(DetallePago.lote_id == lote_id)
        .order_by(DetallePago.id.asc())
        .all()
    )

    if not items:
        raise HTTPException(status_code=404, detail="No hay registros para este lote")

    lines = []

    # HEADER
    lines.append("GrupoPago,TipoPago,id_primario,id_alterno,cuenta,monto,programa")

    for it in items:

        if not it.numero_cuenta:
            raise HTTPException(
                status_code=400,
                detail=f"El expediente {it.expediente_id} no tiene cuenta bancaria"
            )

        line = ",".join([
            str(lote_id),                       # 🔥 GrupoPago
            "162078",                           # TipoPago (quemado)
            str(it.expediente_id),              # id_primario
            str(it.expediente_id),              # id_alterno
            it.numero_cuenta,                   # cuenta
            f"{float(it.monto_asignado):.2f}",  # monto
            "2",                                # programa (quemado)
        ])

        lines.append(line)

    content = "\n".join(lines)

    buffer = io.BytesIO(content.encode("utf-8"))

    return StreamingResponse(
        buffer,
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=detalle_lote_{lote_id}.csv"
        }
    )

def eliminar_lote_pago(
    db: Session,
    lote_id: int,
    eliminado_por: str | None = None,
):
    lote = db.execute(
        text("""
            SELECT id, estado, observacion
            FROM lote_pago
            WHERE id = :lote_id
        """),
        {"lote_id": lote_id}
    ).mappings().first()

    if not lote:
        raise HTTPException(
            status_code=404,
            detail="La planilla no existe"
        )

    if lote["estado"] == "ELIMINADO":
        raise HTTPException(
            status_code=400,
            detail="La planilla ya fue eliminada"
        )

    observacion_actual = lote["observacion"] or ""

    texto_eliminacion = "Planilla marcada como ELIMINADO"
    if eliminado_por:
        texto_eliminacion += f" por {eliminado_por}"

    nueva_observacion = (
        f"{observacion_actual} | {texto_eliminacion}"
        if observacion_actual.strip()
        else texto_eliminacion
    )

    db.execute(
        text("""
            UPDATE lote_pago
            SET
                estado = 'ELIMINADO',
                observacion = :observacion
            WHERE id = :lote_id
        """),
        {
            "lote_id": lote_id,
            "observacion": nueva_observacion,
        }
    )

    db.commit()

    return {
        "ok": True,
        "message": "Planilla eliminada correctamente",
        "id": lote_id,
        "estado": "ELIMINADO",
    }